"""패널 DataFrame → xlsx (panel + README 시트) 공통 라이터."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CENTER = Alignment(horizontal="center", vertical="center")


def write_panel_xlsx(
    panel: pd.DataFrame,
    out: Path,
    *,
    headers_kr: dict[str, str],
    column_order: list[str],
    formula_columns: dict[str, str] | None = None,
    readme_rows: list[tuple[str, str]] | None = None,
    column_widths: list[int] | None = None,
) -> None:
    """
    panel        : long-format DataFrame (sido_code, sido, year, ...)
    headers_kr   : 컬럼명 → 한글 헤더
    column_order : 출력 컬럼 순서
    formula_columns : 컬럼명 → 엑셀 수식 템플릿 ('{r}' 자리표시자)
                      예: {"total": "=IFERROR(D{r}+E{r},\"\")"}
    readme_rows  : README 시트 (key, value) 행
    """
    formula_columns = formula_columns or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "panel"

    for j, col in enumerate(column_order, start=1):
        c = ws.cell(row=1, column=j, value=headers_kr.get(col, col))
        c.fill, c.font, c.alignment = _HEADER_FILL, _HEADER_FONT, _CENTER

    for i, row in enumerate(panel[column_order].itertuples(index=False), start=2):
        for j, col in enumerate(column_order, start=1):
            val = getattr(row, col)
            if col in formula_columns:
                ws.cell(row=i, column=j, value=formula_columns[col].format(r=i))\
                    .number_format = "#,##0"
            elif pd.isna(val):
                continue
            elif col in {"sido_code", "year"}:
                cell = ws.cell(row=i, column=j, value=int(val) if col == "year" else str(val))
                cell.alignment = _CENTER
            elif isinstance(val, str):
                ws.cell(row=i, column=j, value=val)
            else:
                ws.cell(row=i, column=j, value=float(val)).number_format = "#,##0"

    widths = column_widths or [12] * len(column_order)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

    if readme_rows:
        meta = wb.create_sheet("README")
        meta.cell(row=1, column=1, value="항목").fill = _HEADER_FILL
        meta.cell(row=1, column=1).font = _HEADER_FONT
        meta.cell(row=1, column=2, value="내용").fill = _HEADER_FILL
        meta.cell(row=1, column=2).font = _HEADER_FONT
        for r, (k, v) in enumerate(readme_rows, start=2):
            meta.cell(row=r, column=1, value=k)
            meta.cell(row=r, column=2, value=v).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        meta.column_dimensions["A"].width = 28
        meta.column_dimensions["B"].width = 90

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
